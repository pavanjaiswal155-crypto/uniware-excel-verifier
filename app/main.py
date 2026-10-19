
from __future__ import annotations

import io
import os
import uuid
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

import pandas as pd
from fastapi import FastAPI, File, Form, HTTPException, UploadFile
from fastapi.responses import FileResponse, HTMLResponse, JSONResponse, StreamingResponse
from fastapi.staticfiles import StaticFiles
from openpyxl import load_workbook
from openpyxl.styles import PatternFill, Font
from pydantic import BaseModel, Field

APP_DIR = Path(__file__).resolve().parent
STATIC_DIR = APP_DIR / "static"
DATA_DIR = Path(os.getenv("DATA_DIR", APP_DIR.parent / "data"))
DATA_DIR.mkdir(parents=True, exist_ok=True)
MAX_UPLOAD_MB = int(os.getenv("MAX_UPLOAD_MB", "20"))
ALLOWED_EXTENSIONS = {".xlsx", ".xls", ".csv"}

app = FastAPI(title="Uniware Excel Verifier", version="0.1.0")
app.mount("/static", StaticFiles(directory=str(STATIC_DIR)), name="static")

# This in-memory registry is appropriate for a single-user local MVP.
# For multi-user production deployment, replace with authenticated storage/database.
SESSIONS: dict[str, dict[str, Any]] = {}


def json_safe(value: Any) -> Any:
    if value is None:
        return None
    try:
        if pd.isna(value):
            return None
    except (TypeError, ValueError):
        pass
    if isinstance(value, (pd.Timestamp, datetime)):
        return value.isoformat()
    if hasattr(value, "item"):
        try:
            return value.item()
        except Exception:
            pass
    if isinstance(value, (str, int, float, bool)):
        return value
    return str(value)


def read_table(path: Path, sheet_name: str | None = None) -> pd.DataFrame:
    suffix = path.suffix.lower()
    try:
        if suffix == ".csv":
            # utf-8-sig handles common Excel-exported CSV files.
            return pd.read_csv(path, dtype=object, encoding="utf-8-sig", keep_default_na=True)
        if suffix == ".xlsx":
            return pd.read_excel(path, sheet_name=sheet_name or 0, dtype=object, engine="openpyxl")
        if suffix == ".xls":
            return pd.read_excel(path, sheet_name=sheet_name or 0, dtype=object, engine="xlrd")
    except Exception as exc:
        raise HTTPException(status_code=400, detail=f"Could not read this file: {exc}") from exc
    raise HTTPException(status_code=400, detail="Unsupported file type. Upload .xlsx, .xls, or .csv.")


def workbook_sheets(path: Path) -> list[str]:
    if path.suffix.lower() == ".csv":
        return ["CSV"]
    try:
        if path.suffix.lower() == ".xlsx":
            from openpyxl import load_workbook
            wb = load_workbook(path, read_only=True, data_only=False)
            names = list(wb.sheetnames)
            wb.close()
            return names
        if path.suffix.lower() == ".xls":
            xls = pd.ExcelFile(path, engine="xlrd")
            return list(xls.sheet_names)
    except Exception as exc:
        raise HTTPException(status_code=400, detail=f"Could not inspect worksheets: {exc}") from exc
    return []


def normalize_header(value: Any) -> str:
    return str(value).strip() if value is not None else ""


def run_validation(df: pd.DataFrame, required_columns: list[str] | None = None,
                   unique_columns: list[str] | None = None) -> dict[str, Any]:
    required_columns = required_columns or []
    unique_columns = unique_columns or []
    issues: list[dict[str, Any]] = []
    headers = [normalize_header(c) for c in df.columns]
    normalized_to_actual: dict[str, str] = {}
    for h in headers:
        normalized_to_actual.setdefault(h.casefold(), h)

    # Duplicate/blank headers are structural problems; use source column labels as seen.
    seen: set[str] = set()
    for col_idx, header in enumerate(headers, start=1):
        if not header:
            issues.append({
                "row": 1, "column": f"Column {col_idx}", "original_value": None,
                "category": "structure", "severity": "error",
                "issue": "Blank column header",
                "explanation": "This column has no header. Add a meaningful column name before importing.",
                "suggestion": None, "status": "open", "rule": "non_blank_header"
            })
        elif header.casefold() in seen:
            issues.append({
                "row": 1, "column": header, "original_value": header,
                "category": "structure", "severity": "error",
                "issue": "Duplicate column header",
                "explanation": "Two or more columns have the same header. Rename one so each column is unambiguous.",
                "suggestion": None, "status": "open", "rule": "unique_headers"
            })
        seen.add(header.casefold())

    for required in required_columns:
        if required.casefold() not in normalized_to_actual:
            issues.append({
                "row": 1, "column": required, "original_value": None,
                "category": "required_column", "severity": "error",
                "issue": "Required column missing",
                "explanation": f"The configured required column '{required}' is not present in this file.",
                "suggestion": f"Add the '{required}' column using the official template.",
                "status": "open", "rule": "required_columns"
            })

    for col in df.columns:
        col_name = normalize_header(col)
        series = df[col]
        # Required-cell check applies only to configured required columns.
        if col_name.casefold() in {x.casefold() for x in required_columns}:
            for idx, val in series.items():
                if pd.isna(val) or (isinstance(val, str) and not val.strip()):
                    issues.append({
                        "row": int(idx) + 2, "column": col_name, "original_value": json_safe(val),
                        "category": "missing_value", "severity": "error",
                        "issue": "Required value is blank",
                        "explanation": f"'{col_name}' is configured as mandatory but this row has no value.",
                        "suggestion": "Enter the correct value from the source system; do not guess identifiers.",
                        "status": "open", "rule": "required_cells"
                    })

    # Exact duplicate data rows (not counting header).
    if len(df) > 0:
        duplicate_mask = df.duplicated(keep="first")
        for idx in df.index[duplicate_mask]:
            issues.append({
                "row": int(idx) + 2, "column": "(entire row)", "original_value": None,
                "category": "duplicate_row", "severity": "warning",
                "issue": "Duplicate row",
                "explanation": "This row exactly matches an earlier data row. Confirm whether duplicates are allowed for this file type.",
                "suggestion": "Compare with the earlier row and remove only if you confirm it is unintended.",
                "status": "open", "rule": "duplicate_rows"
            })

    for requested_col in unique_columns:
        actual = normalized_to_actual.get(requested_col.casefold())
        if actual is None:
            issues.append({
                "row": 1, "column": requested_col, "original_value": None,
                "category": "unique_identifier", "severity": "warning",
                "issue": "Uniqueness check skipped",
                "explanation": f"Configured unique column '{requested_col}' is missing, so duplicate checking could not run.",
                "suggestion": "Confirm the column name against the official template.",
                "status": "open", "rule": "unique_columns"
            })
            continue
        nonblank = df[actual].notna() & df[actual].astype(str).str.strip().ne("")
        duplicated = df[actual].astype(str).str.strip().duplicated(keep=False) & nonblank
        for idx in df.index[duplicated]:
            val = json_safe(df.loc[idx, actual])
            issues.append({
                "row": int(idx) + 2, "column": normalize_header(actual), "original_value": val,
                "category": "duplicate_identifier", "severity": "error",
                "issue": f"Duplicate value in {actual}",
                "explanation": f"The value '{val}' appears more than once in the configured unique column. Confirm the business rule before using this check.",
                "suggestion": "Check the source data and resolve the duplicate using the authoritative system.",
                "status": "open", "rule": "unique_columns"
            })

    # Numeric-looking columns: flag values that mix numeric and nonnumeric data only as warnings.
    for col in df.columns:
        name = normalize_header(col)
        lower = name.casefold()
        if any(token in lower for token in ("qty", "quantity", "price", "amount", "weight", "stock")):
            vals = df[col]
            for idx, val in vals.items():
                if pd.isna(val) or (isinstance(val, str) and not val.strip()):
                    continue
                try:
                    num = float(str(val).replace(",", "").strip())
                    if not pd.notna(num):
                        raise ValueError()
                except (ValueError, TypeError):
                    issues.append({
                        "row": int(idx) + 2, "column": name, "original_value": json_safe(val),
                        "category": "numeric_format", "severity": "warning",
                        "issue": "Value may not be numeric",
                        "explanation": f"'{name}' looks like a numeric field based on its header, but this value could not be parsed as a number.",
                        "suggestion": "Check the value and the configured format for this file type.",
                        "status": "open", "rule": "numeric_heuristic"
                    })

    error_count = sum(1 for i in issues if i["severity"] == "error")
    warning_count = sum(1 for i in issues if i["severity"] == "warning")
    issue_rows = {i["row"] for i in issues if i["row"] >= 2}
    error_rows = {i["row"] for i in issues if i["row"] >= 2 and i["severity"] == "error"}
    warning_rows = {i["row"] for i in issues if i["row"] >= 2 and i["severity"] == "warning"}
    data_rows = len(df)
    # Structural/header errors are blocking; row-based errors are also blocking.
    if error_count:
        status = "FAIL"
    elif warning_count:
        status = "PASS WITH WARNINGS"
    else:
        status = "PASS"
    return {
        "status": status,
        "data_rows_checked": data_rows,
        "passed_rows": max(0, data_rows - len(error_rows)),
        "failed_rows": len(error_rows),
        "rows_with_warnings": len(warning_rows),
        "error_count": error_count,
        "warning_count": warning_count,
        "issues": issues,
        "verification_scope": "Generic data-quality checks plus only the configured rules. This is not proof of Uniware acceptance.",
        "required_columns": required_columns,
        "unique_columns": unique_columns,
        "verified_at": datetime.now(timezone.utc).isoformat()
    }


@app.get("/", response_class=HTMLResponse)
def index():
    return (STATIC_DIR / "index.html").read_text(encoding="utf-8")


@app.get("/api/health")
def health():
    return {"status": "ok", "version": app.version}


@app.post("/api/upload")
async def upload_file(file: UploadFile = File(...)):
    filename = Path(file.filename or "upload").name
    ext = Path(filename).suffix.lower()
    if ext not in ALLOWED_EXTENSIONS:
        raise HTTPException(status_code=415, detail="Unsupported file type. Upload .xlsx, .xls, or .csv.")
    content = await file.read((MAX_UPLOAD_MB * 1024 * 1024) + 1)
    if not content:
        raise HTTPException(status_code=400, detail="The uploaded file is empty.")
    if len(content) > MAX_UPLOAD_MB * 1024 * 1024:
        raise HTTPException(status_code=413, detail=f"File exceeds the {MAX_UPLOAD_MB} MB limit.")
    if ext == ".xlsx" and not content.startswith(b"PK"):
        raise HTTPException(status_code=400, detail="This does not look like a valid .xlsx workbook.")
    session_id = uuid.uuid4().hex
    path = DATA_DIR / f"{session_id}{ext}"
    path.write_bytes(content)
    sheets = workbook_sheets(path)
    df = read_table(path, sheets[0] if ext != ".csv" else None)
    # Keep raw uploaded bytes and selected sheet for later export.
    SESSIONS[session_id] = {
        "path": path, "filename": filename, "extension": ext,
        "sheets": sheets, "sheet": sheets[0] if sheets else None,
        "required_columns": [], "unique_columns": [],
        "df": df, "original_df": df.copy(deep=True), "approved_changes": [],
        "result": None
    }
    return {
        "session_id": session_id, "filename": filename, "extension": ext,
        "sheets": sheets, "selected_sheet": sheets[0] if sheets else None,
        "rows": len(df), "columns": len(df.columns),
        "headers": [normalize_header(c) for c in df.columns],
        "preview": [[json_safe(v) for v in row] for row in df.head(100).itertuples(index=False, name=None)],
        "preview_limit": 100
    }


class VerifyRequest(BaseModel):
    session_id: str
    sheet: str | None = None
    required_columns: list[str] = Field(default_factory=list)
    unique_columns: list[str] = Field(default_factory=list)


@app.post("/api/verify")
def verify(req: VerifyRequest):
    session = SESSIONS.get(req.session_id)
    if not session:
        raise HTTPException(status_code=404, detail="Upload session not found. Please upload the file again.")
    if req.sheet and req.sheet not in session["sheets"]:
        raise HTTPException(status_code=400, detail="Selected worksheet does not exist.")
    if session["extension"] != ".csv" and req.sheet:
        session["df"] = read_table(session["path"], req.sheet)
        session["sheet"] = req.sheet
    session["required_columns"] = [str(x).strip() for x in req.required_columns if str(x).strip()]
    session["unique_columns"] = [str(x).strip() for x in req.unique_columns if str(x).strip()]
    result = run_validation(session["df"], session["required_columns"], session["unique_columns"])
    result["filename"] = session["filename"]
    result["sheet"] = session["sheet"]
    result["columns"] = [normalize_header(c) for c in session["df"].columns]
    session["result"] = result
    return result


class Correction(BaseModel):
    session_id: str
    row: int = Field(ge=2)
    column: str
    new_value: str
    approved: bool


class ApplyCorrections(BaseModel):
    session_id: str
    corrections: list[Correction]


@app.post("/api/corrections")
def apply_corrections(req: ApplyCorrections):
    session = SESSIONS.get(req.session_id)
    if not session:
        raise HTTPException(status_code=404, detail="Upload session not found.")
    df = session["df"].copy(deep=True)
    headers = {normalize_header(c): c for c in df.columns}
    applied = []
    for correction in req.corrections:
        if correction.session_id != req.session_id:
            raise HTTPException(status_code=400, detail="Correction session mismatch.")
        if not correction.approved:
            continue
        if correction.column not in headers:
            raise HTTPException(status_code=400, detail=f"Unknown column: {correction.column}")
        idx = correction.row - 2
        if idx < 0 or idx >= len(df):
            raise HTTPException(status_code=400, detail=f"Row {correction.row} is outside the data range.")
        col = headers[correction.column]
        old = json_safe(df.iloc[idx][col])
        df.at[df.index[idx], col] = correction.new_value
        applied.append({"row": correction.row, "column": correction.column,
                        "original_value": old, "new_value": correction.new_value})
    session["df"] = df
    session["approved_changes"].extend(applied)
    result = run_validation(df, session["required_columns"], session["unique_columns"])
    result["filename"] = session["filename"]
    result["sheet"] = session["sheet"]
    result["columns"] = [normalize_header(c) for c in df.columns]
    session["result"] = result
    return {"applied": applied, "applied_count": len(applied), "result": result}


def make_report_csv(session: dict[str, Any]) -> bytes:
    result = session.get("result")
    if not result:
        raise HTTPException(status_code=409, detail="Run verification before downloading a report.")
    rows = []
    for issue in result["issues"]:
        rows.append({
            "file": session["filename"], "sheet": session["sheet"],
            "verification_status": result["status"], "row": issue["row"],
            "column": issue["column"], "original_value": issue["original_value"],
            "category": issue["category"], "severity": issue["severity"],
            "issue": issue["issue"], "explanation": issue["explanation"],
            "suggestion": issue["suggestion"], "review_status": issue["status"],
            "rule": issue["rule"]
        })
    if not rows:
        rows = [{"file": session["filename"], "sheet": session["sheet"],
                 "verification_status": result["status"], "issue": "No issues found"}]
    return pd.DataFrame(rows).to_csv(index=False).encode("utf-8-sig")


@app.get("/api/report/{session_id}.csv")
def download_report_csv(session_id: str):
    session = SESSIONS.get(session_id)
    if not session:
        raise HTTPException(status_code=404, detail="Upload session not found.")
    content = make_report_csv(session)
    return StreamingResponse(io.BytesIO(content), media_type="text/csv",
        headers={"Content-Disposition": f'attachment; filename="verification-report-{session_id[:8]}.csv"'})


@app.get("/api/report/{session_id}.xlsx")
def download_report_xlsx(session_id: str):
    session = SESSIONS.get(session_id)
    if not session:
        raise HTTPException(status_code=404, detail="Upload session not found.")
    if not session.get("result"):
        raise HTTPException(status_code=409, detail="Run verification before downloading a report.")
    result = session["result"]
    summary = pd.DataFrame([{
        "File": session["filename"], "Worksheet": session["sheet"],
        "Status": result["status"], "Rows checked": result["data_rows_checked"],
        "Passed rows": result["passed_rows"], "Failed rows": result["failed_rows"],
        "Rows with warnings": result["rows_with_warnings"],
        "Errors": result["error_count"], "Warnings": result["warning_count"],
        "Scope": result["verification_scope"]
    }])
    issues = pd.DataFrame([{
        "Row": i["row"], "Column": i["column"], "Original value": i["original_value"],
        "Category": i["category"], "Severity": i["severity"], "Issue": i["issue"],
        "Explanation": i["explanation"], "Suggested correction": i["suggestion"],
        "Review status": i["status"], "Rule": i["rule"]
    } for i in result["issues"]])
    if issues.empty:
        issues = pd.DataFrame([{"Issue": "No issues found"}])
    output = io.BytesIO()
    with pd.ExcelWriter(output, engine="openpyxl") as writer:
        summary.to_excel(writer, index=False, sheet_name="Summary")
        issues.to_excel(writer, index=False, sheet_name="Issues")
        changes = pd.DataFrame(session["approved_changes"])
        if changes.empty:
            changes = pd.DataFrame([{"message": "No approved changes applied"}])
        changes.to_excel(writer, index=False, sheet_name="Approved Changes")
    output.seek(0)
    return StreamingResponse(output, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="verification-report-{session_id[:8]}.xlsx"'})


@app.get("/api/corrected/{session_id}")
def download_corrected(session_id: str):
    session = SESSIONS.get(session_id)
    if not session:
        raise HTTPException(status_code=404, detail="Upload session not found.")
    if not session.get("result"):
        raise HTTPException(status_code=409, detail="Run verification before exporting.")
    df = session["df"]
    out = io.BytesIO()
    if session["extension"] == ".csv":
        out.write(df.to_csv(index=False).encode("utf-8-sig"))
        out.seek(0)
        media = "text/csv"
        suffix = ".csv"
    else:
        # Export the selected sheet as a clean workbook; original formatting and other sheets
        # are not preserved in this MVP, so make this limitation explicit in the UI/README.
        with pd.ExcelWriter(out, engine="openpyxl") as writer:
            df.to_excel(writer, index=False, sheet_name=(session["sheet"] or "Verified")[:31])
        out.seek(0)
        media = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        suffix = ".xlsx"
    return StreamingResponse(out, media_type=media,
        headers={"Content-Disposition": f'attachment; filename="verified-{Path(session["filename"]).stem}{suffix}"'})


@app.delete("/api/session/{session_id}")
def delete_session(session_id: str):
    session = SESSIONS.pop(session_id, None)
    if not session:
        raise HTTPException(status_code=404, detail="Session not found.")
    try:
        session["path"].unlink(missing_ok=True)
    except Exception:
        pass
    return {"deleted": True}
